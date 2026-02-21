"""
Парсер каталога bradytool.ru -> Excel.

Собирает товары из всех категорий интернет-магазина Brady,
обогащает данные со страниц карточек и сохраняет в .xlsx.
"""

from __future__ import annotations

import argparse
import logging
import math
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field, fields, asdict
from pathlib import Path
from typing import Sequence

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# ---------------------------------------------------------------------------
# Конфигурация
# ---------------------------------------------------------------------------

BASE_URL = "https://bradytool.ru"

CATEGORIES: list[tuple[str, str]] = [
    ("printerny", "Принтеры"),
    ("kartridzhi-i-lenty", "Картриджи и ленты"),
    ("blokiruyushchie-ustrojstva", "Блокирующие устройства"),
    ("ehtiketki-i-markery-brdy", "Этикетки и маркеры"),
    ("software", "Софт"),
    ("gotovaya-markirovka", "Готовая маркировка"),
    ("ribbony", "Риббоны"),
    ("aksessuary", "Аксессуары"),
]

DEFAULT_WORKERS = 5
DEFAULT_TIMEOUT = 30
DEFAULT_LIMIT = 100
DEFAULT_SAVE_EVERY = 200
DEFAULT_OUTPUT = "bradytool_catalog.xlsx"
DEFAULT_RETRIES = 3
DELAY_BETWEEN_PAGES = 0.3

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0.0.0 Safari/537.36"
)

log = logging.getLogger("brady_parser")

# ---------------------------------------------------------------------------
# Модель данных
# ---------------------------------------------------------------------------

EXCEL_COLUMNS: list[str] = [
    "Ссылка",
    "Раздел",
    "Артикул",
    "Товар",
    "Описание",
    "Цена",
    "Наличие",
    "Производитель",
    "Фото",
]


@dataclass
class Product:
    url: str = ""
    category: str = ""
    sku: str = ""
    name: str = ""
    description: str = ""
    price: str = ""
    availability: str = ""
    manufacturer: str = ""
    image: str = ""

    def to_row(self) -> dict[str, str]:
        return {
            "Ссылка": self.url,
            "Раздел": self.category,
            "Артикул": self.sku,
            "Товар": self.name,
            "Описание": self.description,
            "Цена": self.price,
            "Наличие": self.availability,
            "Производитель": self.manufacturer,
            "Фото": self.image,
        }


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------


def create_session() -> requests.Session:
    session = requests.Session()
    session.headers.update({
        "User-Agent": USER_AGENT,
        "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.3",
    })
    return session


def fetch(
    session: requests.Session,
    url: str,
    *,
    retries: int = DEFAULT_RETRIES,
    timeout: int = DEFAULT_TIMEOUT,
) -> requests.Response | None:
    for attempt in range(retries):
        try:
            resp = session.get(url, timeout=timeout)
            resp.raise_for_status()
            return resp
        except requests.RequestException as exc:
            if attempt < retries - 1:
                delay = 2 * (attempt + 1)
                log.debug("Retry %d/%d for %s (sleep %ds): %s", attempt + 1, retries, url, delay, exc)
                time.sleep(delay)
            else:
                log.warning("Failed after %d attempts: %s — %s", retries, url, exc)
    return None


# ---------------------------------------------------------------------------
# Вспомогательные функции
# ---------------------------------------------------------------------------

_CACHE_DIM_RE = re.compile(r"-\d+x\d+(\.\w+)$")


def original_image_url(cached_url: str) -> str:
    """Преобразует URL кешированной миниатюры OpenCart в URL оригинала."""
    return _CACHE_DIM_RE.sub(r"\1", cached_url.replace("/cache", ""))


def get_total_pages(soup: BeautifulSoup, items_per_page: int) -> int:
    max_page = 1

    for link in soup.select(".pagination li a"):
        m = re.search(r"page=(\d+)", link.get("href", ""))
        if m:
            max_page = max(max_page, int(m.group(1)))

    text = soup.get_text(" ", strip=True)

    m = re.search(r"всего\s+(\d+)\s+страниц", text, re.IGNORECASE)
    if m:
        max_page = max(max_page, int(m.group(1)))

    m = re.search(r"из\s+(\d+)\s*\(", text)
    if m:
        total_items = int(m.group(1))
        max_page = max(max_page, math.ceil(total_items / items_per_page))

    return max_page


# ---------------------------------------------------------------------------
# Парсинг страниц
# ---------------------------------------------------------------------------


def parse_listing_page(html: bytes, category_name: str) -> list[Product]:
    """Извлекает карточки товаров со страницы каталога."""
    soup = BeautifulSoup(html, "lxml")
    products: list[Product] = []

    for card in soup.select(".product-card"):
        p = Product(category=category_name)

        link_el = card.select_one(".image a")
        if link_el:
            p.url = link_el.get("href", "")
            img_el = link_el.select_one("img")
            if img_el:
                p.image = original_image_url(img_el.get("src", ""))

        name_el = card.select_one(".product-name a, h4 a")
        if name_el:
            p.name = name_el.get_text(strip=True)
            if not p.url:
                p.url = name_el.get("href", "")

        desc_el = card.select_one(".description")
        if desc_el:
            p.description = desc_el.get_text(strip=True)

        price_el = card.select_one(".price")
        if price_el:
            p.price = re.sub(r"\s+", " ", price_el.get_text(strip=True))

        if p.url:
            products.append(p)

    return products


def enrich_from_product_page(session: requests.Session, product: Product) -> Product:
    """Дополняет данные товара информацией с его страницы."""
    if not product.url:
        return product

    resp = fetch(session, product.url)
    if resp is None:
        return product

    soup = BeautifulSoup(resp.content, "lxml")

    h1 = soup.select_one("h1.heading-title")
    if h1:
        product.name = h1.get_text(strip=True)

    for li in soup.select(".list-unstyled li"):
        text = li.get_text(strip=True)
        if "Код товара:" in text:
            product.sku = text.split("Код товара:")[-1].strip()
        elif "Наличие:" in text:
            product.availability = text.split("Наличие:")[-1].strip()
        elif "Производитель:" in text:
            product.manufacturer = text.split("Производитель:")[-1].strip()

    price_el = soup.select_one(".product-right-block h2")
    if price_el:
        product.price = price_el.get_text(strip=True)

    desc_el = soup.select_one("#tab-description")
    if desc_el:
        product.description = desc_el.get_text(strip=True)

    thumb_link = soup.select_one("a.thumbnail")
    if thumb_link:
        href = thumb_link.get("href", "")
        if href:
            product.image = original_image_url(href)

    return product


# ---------------------------------------------------------------------------
# Excel-экспорт
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
COL_WIDTHS = [50, 25, 20, 60, 80, 15, 18, 15, 70]


def save_to_excel(products: Sequence[Product], path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Каталог Brady"

    for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    for row_idx, product in enumerate(products, 2):
        row = product.to_row()
        for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))

    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(str(path))


# ---------------------------------------------------------------------------
# Основные этапы
# ---------------------------------------------------------------------------


def collect_products(
    session: requests.Session,
    categories: list[tuple[str, str]],
    items_per_page: int = DEFAULT_LIMIT,
) -> list[Product]:
    """Этап 1: обход страниц каталога, сбор предварительных данных."""
    all_products: list[Product] = []

    for slug, name in categories:
        log.info("Категория: %s", name)
        url = f"{BASE_URL}/{slug}?limit={items_per_page}"
        resp = fetch(session, url)
        if resp is None:
            log.error("Не удалось загрузить категорию «%s»", name)
            continue

        soup = BeautifulSoup(resp.content, "lxml")
        total_pages = get_total_pages(soup, items_per_page)
        log.info("  Страниц: %d", total_pages)

        page_products = parse_listing_page(resp.content, name)
        all_products.extend(page_products)
        log.info("  Стр. 1/%d — %d товаров", total_pages, len(page_products))

        for page_num in range(2, total_pages + 1):
            page_url = f"{BASE_URL}/{slug}?limit={items_per_page}&page={page_num}"
            resp = fetch(session, page_url)
            if resp is None:
                continue
            page_products = parse_listing_page(resp.content, name)
            all_products.extend(page_products)
            log.info("  Стр. %d/%d — %d товаров", page_num, total_pages, len(page_products))
            time.sleep(DELAY_BETWEEN_PAGES)

    return all_products


def enrich_products(
    products: list[Product],
    *,
    workers: int = DEFAULT_WORKERS,
    save_every: int = DEFAULT_SAVE_EVERY,
    output_path: str | Path = DEFAULT_OUTPUT,
) -> list[Product]:
    """Этап 2: параллельная загрузка карточек товаров."""
    total = len(products)
    if total == 0:
        return products

    enriched: list[Product | None] = [None] * total
    counter = 0
    lock = threading.Lock()
    stop_event = threading.Event()

    def _worker(idx: int, product: Product) -> tuple[int, Product]:
        if stop_event.is_set():
            return idx, product
        session = create_session()
        try:
            return idx, enrich_from_product_page(session, product)
        except Exception as exc:
            log.warning("Ошибка при обработке %s: %s", product.url, exc)
            return idx, product

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {
            pool.submit(_worker, i, p): i for i, p in enumerate(products)
        }

        try:
            for future in as_completed(futures):
                idx, result = future.result()
                with lock:
                    enriched[idx] = result
                    counter += 1
                    current = counter

                if current % 50 == 0 or current == total:
                    log.info("  [%d/%d] (%.1f%%)", current, total, current / total * 100)

                if current % save_every == 0:
                    with lock:
                        partial = [p for p in enriched if p is not None]
                    save_to_excel(partial, output_path)
                    log.info("  -- Промежуточное сохранение (%d записей)", len(partial))

        except KeyboardInterrupt:
            log.warning("Прервано пользователем — сохраняем собранные данные…")
            stop_event.set()

    return [p for p in enriched if p is not None]


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Парсер каталога bradytool.ru -> Excel",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Примеры:\n"
            "  python parser_brady.py                       # полный парсинг\n"
            "  python parser_brady.py --test                # тест на 10 товарах\n"
            "  python parser_brady.py -w 10 -o result.xlsx  # 10 потоков, свой файл\n"
            '  python parser_brady.py -c "Принтеры"         # только одна категория\n'
        ),
    )
    p.add_argument(
        "--test", action="store_true",
        help="тестовый режим — обработать только 10 товаров",
    )
    p.add_argument(
        "-o", "--output", default=DEFAULT_OUTPUT,
        help=f"имя выходного Excel-файла (по умолчанию: {DEFAULT_OUTPUT})",
    )
    p.add_argument(
        "-w", "--workers", type=int, default=DEFAULT_WORKERS,
        help=f"количество потоков загрузки (по умолчанию: {DEFAULT_WORKERS})",
    )
    p.add_argument(
        "-c", "--category", action="append", default=None,
        help="фильтр по названию или slug категории (можно указать несколько раз)",
    )
    p.add_argument(
        "--save-every", type=int, default=DEFAULT_SAVE_EVERY,
        help=f"промежуточное сохранение каждые N товаров (по умолчанию: {DEFAULT_SAVE_EVERY})",
    )
    p.add_argument(
        "-v", "--verbose", action="store_true",
        help="подробный вывод (уровень DEBUG)",
    )
    return p


def setup_logging(verbose: bool = False) -> None:
    level = logging.DEBUG if verbose else logging.INFO
    fmt = "%(asctime)s  %(levelname)-8s  %(message)s"
    logging.basicConfig(level=level, format=fmt, datefmt="%H:%M:%S", stream=sys.stdout)


def main(argv: list[str] | None = None) -> None:
    args = build_parser().parse_args(argv)
    setup_logging(args.verbose)

    log.info("=" * 50)
    log.info("Парсер каталога bradytool.ru")
    log.info("=" * 50)

    categories = CATEGORIES
    if args.category:
        allowed = {c.lower() for c in args.category}
        categories = [
            (s, n) for s, n in CATEGORIES
            if n.lower() in allowed or s.lower() in allowed
        ]
        if not categories:
            log.error(
                "Категории не найдены. Доступные: %s",
                ", ".join(f"{n} ({s})" for s, n in CATEGORIES),
            )
            sys.exit(1)
        log.info("Выбранные категории: %s", ", ".join(n for _, n in categories))

    session = create_session()

    log.info("[1/3] Сбор ссылок на товары из каталога…")
    products = collect_products(session, categories)
    log.info("Всего найдено товаров: %d", len(products))

    if args.test:
        products = products[:10]
        log.info("[ТЕСТ] Ограничено до %d товаров", len(products))

    log.info("[2/3] Загрузка карточек товаров (%d шт.)…", len(products))
    products = enrich_products(
        products,
        workers=args.workers,
        save_every=args.save_every,
        output_path=args.output,
    )

    log.info("[3/3] Сохранение в %s…", args.output)
    save_to_excel(products, args.output)
    log.info("Готово! Сохранено %d товаров -> %s", len(products), args.output)


if __name__ == "__main__":
    main()
